"""Plan exports: Excel workbook, one-page PDF, JSON per market (story F5)."""

from __future__ import annotations

import io
import json
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

HEADER = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="5B5BF7")


def _sheet(wb: Workbook, title: str, header: list[str], rows: list[list[Any]]) -> None:
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font, cell.fill = HEADER, FILL
    for r in rows:
        ws.append(r)
    for col in ws.columns:
        width = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, width + 2))


def _concept(group: dict[str, Any], cid: str | None) -> dict[str, Any] | None:
    return next((c for c in group.get("concepts", []) if c["id"] == cid), None)


def xlsx(campaign: dict[str, Any]) -> bytes:
    result, plan = campaign["result"], campaign["result"]["plan"]
    groups = {g["key"]: g for g in result["groups"]}
    wb = Workbook()
    wb.remove(wb.active)
    cur = plan["currency"]
    # First sheet: what is data and what is AI (EU AI Act transparency), before any number.
    _sheet(wb, "Read me", ["About this plan"], [
        [f"Tastegeist campaign plan: {campaign.get('title', '')}"],
        ["Taste data, lift, zones and interest keywords: Qloo, computed by Tastegeist."],
        ["Creative concepts, persona reactions and audience names: written by AI (Claude) from the Qloo data. Review before use."],
        ["Séance results are a synthetic simulation: directional, not a survey. Validate with the A/B tests."],
        ["Imagery referenced in concepts belongs to its owners and is not licensed for advertising."],
    ])
    _sheet(wb, "Budget", ["Zone", "Persona", "Markets", "Share", f"Budget ({cur})"] + [c["label"] for c in plan["channels"]],
           [[z["label"], z["name"], ", ".join(z["markets"]), round(z["share"], 4), z["budget"]]
            + [next((c["budget"] for c in z["channels"] if c["channel"] == ch["channel"]), 0) for ch in plan["channels"]] for z in plan["zones"]])
    _sheet(wb, "Audiences", ["Zone", "Ages", "Gender", "Domain", "Interest keywords (Qloo ledger)"],
           [[z["name"], ", ".join(z["audience"]["age"]), z["audience"]["gender"], d, ", ".join(names)]
            for z in plan["zones"] for d, names in z["audience"]["interests"].items()])
    _sheet(wb, "Geo", ["Zone", "Place", "Latitude", "Longitude", "Radius (km)"],
           [[z["name"], g["name"], g["lat"], g["lon"], g["radius_km"]] for z in plan["zones"] for g in z["geo"]])
    _sheet(wb, "Flighting", ["Phase", "Channel", "Start", "End"], [[f["phase"], f["label"], f["start"], f["end"]] for f in plan["flighting"]])
    rows = []
    for z in plan["zones"]:
        c = _concept(groups[z["key"]], (z.get("creative") or {}).get("id"))
        if c:
            rows.append([z["name"], c.get("direction"), (c.get("headline") or {}).get("text"), (c.get("big_idea") or {}).get("text"), (c.get("body") or {}).get("text")])
    _sheet(wb, "Creatives", ["Zone", "Direction", "Headline", "Big idea", "Body"], rows)
    _sheet(wb, "AB tests", ["Zone", "Variant A", "Variant B", "Split", "Metric", "Impressions per arm", "Assumption"],
           [[z["name"], z["ab_test"]["a"]["headline"], z["ab_test"]["b"]["headline"], z["ab_test"]["split"], z["ab_test"]["metric"],
             z["ab_test"]["impressions_per_arm"], z["ab_test"]["assumption"]] for z in plan["zones"] if z.get("ab_test")])
    _sheet(wb, "Ledger", ["Zone", "Entity", "Domain", "Lift", "Qloo entity id"],
           [[g.get("persona", {}).get("name") or g["label"], e["name"], e["domain"], e["lift"], e["id"]] for g in result["groups"] for e in g["ledger"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _logo(data_url: str | None, height: float):
    """A PNG or JPEG data URL as a reportlab image of the given height (SVG and WebP are skipped)."""
    import base64

    from reportlab.lib.utils import ImageReader
    from reportlab.platypus import Image

    if not data_url or not data_url.startswith(("data:image/png;base64,", "data:image/jpeg;base64,")):
        return None
    try:
        raw = io.BytesIO(base64.b64decode(data_url.split(",", 1)[1]))
        w, h = ImageReader(raw).getSize()
        raw.seek(0)
        return Image(raw, width=height * w / h, height=height)
    except Exception:
        return None


def pdf(campaign: dict[str, Any], branding: dict[str, Any] | None = None) -> bytes:
    """The plan on one landscape page. With a brand kit: the brand color, the agency's logo and a prepared-by line."""
    branding = branding or {}
    result, plan = campaign["result"], campaign["result"]["plan"]
    groups = {g["key"]: g for g in result["groups"]}
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=14 * mm, rightMargin=14 * mm, topMargin=12 * mm, bottomMargin=12 * mm)
    ink = colors.HexColor("#16130F")
    brand = colors.HexColor((branding.get("brand") or {}).get("color") or "#16130F")
    h1 = ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=ink)
    sub = ParagraphStyle("sub", fontName="Helvetica", fontSize=10, leading=13, textColor=colors.HexColor("#64748B"))
    cell = ParagraphStyle("cell", fontName="Helvetica", fontSize=8, leading=10, textColor=ink)
    money = lambda n: f"{plan['currency']} {n:,.0f}"  # noqa: E731

    rows = [["Zone", "Markets", "Budget", "Creative (chosen)", "Interest keywords (Qloo)", "A/B"]]
    for z in plan["zones"]:
        c = _concept(groups[z["key"]], (z.get("creative") or {}).get("id"))
        kw = "; ".join(f"{d}: {', '.join(n[:3])}" for d, n in z["audience"]["interests"].items())
        rows.append([
            Paragraph(f"<b>{z['name']}</b>", cell), Paragraph(", ".join(z["markets"]), cell), money(z["budget"]),
            Paragraph((c.get("headline") or {}).get("text", "") if c else "", cell), Paragraph(kw, cell),
            Paragraph(f"{z['ab_test']['impressions_per_arm']:,} / arm" if z.get("ab_test") else "", cell),
        ])
    table = Table(rows, colWidths=[38 * mm, 38 * mm, 26 * mm, 62 * mm, 82 * mm, 22 * mm], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), brand), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8),
        ("VALIGN", (0, 0), (-1, -1), "TOP"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F1F2F7")]),
        ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#E4E6EE")),
    ]))
    agency = branding.get("agency") or {}
    prepared = " ".join(x for x in [f"Prepared by {agency['name']}" if agency.get("name") else "", f"for {branding['client']}" if branding.get("client") else ""] if x)
    logo = _logo(agency.get("logo"), 10 * mm)
    story = [
        *([logo, Spacer(1, 4 * mm)] if logo else []),
        *([Paragraph(prepared, sub), Spacer(1, 2 * mm)] if prepared else []),
        Paragraph(campaign.get("headline") or campaign["title"], h1),
        Spacer(1, 3 * mm),
        Paragraph(f"{campaign['title']} · {plan['start']} to {plan['end']} · {money(plan['budget'])} · channels: {', '.join(c['label'] for c in plan['channels'])}", sub),
        Spacer(1, 6 * mm),
        table,
        Spacer(1, 5 * mm),
        Paragraph(f"Method: {plan['method']} Every interest keyword is a Qloo entity with lift above 1 in that zone. Concepts and persona reactions are AI-generated (Claude) from Qloo data; séance results are synthetic and directional. Made with Tastegeist.", sub),
    ]
    doc.build(story)
    return buf.getvalue()


def json_markets(campaign: dict[str, Any]) -> bytes:
    plan = campaign["result"]["plan"]
    out = {
        "campaign": campaign["title"],
        "flight": {"start": plan["start"], "end": plan["end"], "currency": plan["currency"]},
        "markets": [
            {"market": m, "zone": z["name"], "budget": round(z["budget"] / max(1, len(z["markets"])), 2), "audience": z["audience"],
             "geo": [g for g in z["geo"] if g["name"] == m] or z["geo"], "creative": z.get("creative"), "channels": z["channels"]}
            for z in plan["zones"] for m in z["markets"]
        ],
    }
    return json.dumps(out, indent=2, ensure_ascii=False).encode()
